from __future__ import annotations
import json
from typing import Any, Dict, List
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PAGE_TITLES = [
 "APPRAISAL NOTE / PROPOSAL GIST",
 "PRICING, CONCESSIONS AND BASIC DATA",
 "BORROWER PROFILE, MANAGEMENT AND SECURITY",
 "BANKING ARRANGEMENT AND BORROWING PROFILE",
 "BUSINESS BACKGROUND AND ACCOUNT CONDUCT",
 "AUDIT, STATUTORY DUES AND COMPLIANCE",
 "EXTERNAL / BUREAU / NEGATIVE LIST VERIFICATION",
 "POLICY BENCHMARKS AND EXPOSURE CHECKS",
 "MULTIPLE BANKING / CONSORTIUM COMPLIANCE",
 "SMA / STRESS / EARLY WARNING REVIEW",
 "OTHER CREDIT PARAMETERS AND MANAGEMENT CHECKS",
 "GROUP / ASSOCIATE EXPOSURE AND JUSTIFICATION",
 "RECOMMENDATION, DEVIATIONS AND CONDITIONS",
 "FINANCIAL PARAMETERS - BALANCE SHEET",
 "FINANCIAL PARAMETERS - PROFITABILITY AND RATIOS",
 "WORKING CAPITAL, CASH FLOW AND OPERATING CYCLE",
 "FINANCIAL PERFORMANCE COMMENTARY",
 "BUSINESS / INDUSTRY / PEER ANALYSIS",
 "COLLATERAL, RISK MITIGATION AND DUE DILIGENCE",
 "FINAL CREDIT VIEW, SOURCE REGISTER AND REVIEW",
]

def _shade(cell, fill='EAF2F8'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def _set_cell_text(cell, value, bold=False, size=8):
    cell.text=''
    p=cell.paragraphs[0]; r=p.add_run('Not available' if value in (None,'') else str(value)); r.bold=bold; r.font.size=Pt(size)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def _table(doc, rows, headers=None, widths=None):
    rows=[list(r) for r in rows]
    if headers: rows=[list(headers)]+rows
    if not rows: return
    cols=max(len(r) for r in rows)
    t=doc.add_table(rows=len(rows),cols=cols); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    try: t.style='Table Grid'
    except: pass
    for i,row in enumerate(rows):
        for j in range(cols):
            _set_cell_text(t.cell(i,j), row[j] if j<len(row) else '', bold=bool(headers and i==0), size=8)
            if headers and i==0: _shade(t.cell(i,j),'D9EAF7')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def _h(doc, text, level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(4); return p

def _page_header(doc, no, title):
    if no>1: doc.add_page_break()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('APPRAISAL NOTE'); r.bold=True; r.font.size=Pt(13)
    p2=doc.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p2.add_run(f'Page {no} - {title}'); r.bold=True; r.font.size=Pt(10)

def _fmt_money(v, unit=''):
    if v in (None,''): return 'Not available'
    try:
        x=float(v)
        return f'{x:,.2f}' + (f' {unit}' if unit else '')
    except: return str(v)

def _field(raw,*keys):
    for k in keys:
        if raw.get(k) not in (None,''): return raw.get(k)
    return None

def _rows_matching(raw, terms, limit=35):
    rows=[]
    for k,v in raw.items():
        if any(t in k.lower() for t in terms): rows.append([k.replace('_',' ').title(),v])
        if len(rows)>=limit: break
    return rows

def build_twenty_page_cam(state: Dict[str,Any], documents: List[Dict[str,Any]], normalized: Dict[str,Any]) -> Document:
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.6); sec.right_margin=Inches(.6)
    styles=doc.styles
    styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(8.5)
    raw=normalized.get('raw_fields') or {}; calc=normalized.get('calculated') or {}; ratios=calc.get('ratios') or {}; ver=normalized.get('verification') or {}
    proposal=normalized.get('proposal') or {}
    unit=raw.get('financial_unit','raw source unit')
    ai=state.get('ai_narratives') or {}

    # 1
    _page_header(doc,1,PAGE_TITLES[0]); _h(doc,'SECTION I: DETAILS OF THE PROPOSAL',1)
    _table(doc,[['Name of Account',proposal.get('company_name')],['CAM ID',state.get('cam_id')],['Proposal','Fresh / Review as captured'],['Facility',proposal.get('loan_type')],['Requested Amount',proposal.get('loan_amount')],['Purpose',proposal.get('loan_purpose')],['Tenure',proposal.get('tenure')],['Repayment',proposal.get('repayment')]],['Particular','Details'])
    doc.add_paragraph(ai.get('executive_summary') or 'Proposal summary is generated from uploaded documents, verified registry data and deterministic calculations. Missing facts are shown as not available.')
    _table(doc,_rows_matching(raw,['facility','limit','outstanding','term_loan','cash_credit','cc_'],25),['Facility Evidence','Source Value'])

    # 2
    _page_header(doc,2,PAGE_TITLES[1]); _table(doc,[['Interest Rate',proposal.get('interest_rate')],['Processing Fee',_field(raw,'processing_fee','processing_charges')],['Pre-closure Charges',_field(raw,'preclosure_charges')],['Asset Classification',_field(raw,'asset_classification')],['Internal Rating',_field(raw,'internal_rating')],['External Rating',_field(raw,'external_rating','credit_rating')]],['Parameter','Value'])
    doc.add_paragraph('Pricing and concessions are taken from the proposal/sanction evidence when available. The system does not invent missing pricing terms.')

    # 3
    _page_header(doc,3,PAGE_TITLES[2]); _table(doc,[['Company Name',proposal.get('company_name')],['CIN',_field(raw,'cin')],['Registered Office',_field(raw,'registered_office')],['Financial Year',_field(raw,'financial_year')],['Managing Director',_field(raw,'managing_director')],['CFO / WTD',_field(raw,'cfo_or_wtd')],['Statutory Auditor',_field(raw,'statutory_auditor')]],['Borrower Field','Extracted Value'])
    _h(doc,'Security / Collateral Evidence',2); _table(doc,_rows_matching(raw,['security','collateral','mortgage','property','guarantee','charge'],30),['Field','Value'])

    # 4
    _page_header(doc,4,PAGE_TITLES[3]); _table(doc,_rows_matching(raw,['bank','borrow','facility','limit','outstanding'],40),['Banking / Borrowing Field','Source Value'])
    doc.add_paragraph('External facility verification is sourced from configured providers where available; POC fallback data is visibly marked as synthetic.')

    # 5
    _page_header(doc,5,PAGE_TITLES[4]); doc.add_paragraph(ai.get('business_overview') or 'Business background is compiled from annual report / corporate documents and not from unrelated generic industry templates.')
    _table(doc,_rows_matching(raw,['business','industry','product','capacity','customer','supplier','sales'],35),['Business Evidence','Value'])
    _h(doc,'Account Conduct',2); add=ver.get('additional_poc_verification') or {}; _table(doc,[[k,v] for k,v in (add.get('banking_conduct') or {}).items()],['Conduct Item','Value'])

    # 6
    _page_header(doc,6,PAGE_TITLES[5]); _table(doc,_rows_matching(raw,['audit','statutory','tax','contingent','litigation','dues'],40),['Audit / Statutory Evidence','Value']); doc.add_paragraph(ai.get('conduct') or 'Compliance observations must be supported by source documents or verification responses.')

    # 7
    _page_header(doc,7,PAGE_TITLES[6]); _table(doc,[[r.get('field'),r.get('document_value'),r.get('verified_value'),r.get('verification_source'),r.get('status')] for r in ver.get('reconciliation',[])],['Field','Document Value','Verified Value','Source','Status']); _h(doc,'POC Bureau / External Checks',2); _table(doc,[[k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v] for k,v in add.items()],['Check','Result'])

    # 8
    _page_header(doc,8,PAGE_TITLES[7]); _table(doc,[['Current Ratio',ratios.get('current_ratio')],['Debt / Equity',ratios.get('debt_equity')],['PAT Margin %',ratios.get('pat_margin_pct')],['Interest Coverage',ratios.get('interest_coverage')]],['Calculated Metric','System Value']); doc.add_paragraph('All ratios on this page are calculated by deterministic Python logic from extracted raw financial values. They are not read from synthetic PDFs or invented by the LLM.')

    # 9
    _page_header(doc,9,PAGE_TITLES[8]); doc.add_paragraph('Consortium / multiple banking compliance is populated only from uploaded sanction letters, bank data, FileSure/MCA charges or configured POC verification data.'); _table(doc,_rows_matching(raw,['consortium','multiple_bank','lender','sanction','bank'],35),['Evidence','Value'])

    # 10
    _page_header(doc,10,PAGE_TITLES[9]); _table(doc,[[k,v] for k,v in (add.get('banking_conduct') or {}).items()],['Stress / Conduct Item','Value']); doc.add_paragraph('SMA / overdue / stress fields are synthetic POC verification only when genuine internal-bank data is unavailable. The source kind is retained in the audit trail.')

    # 11
    _page_header(doc,11,PAGE_TITLES[10]); _table(doc,_rows_matching(raw,['director','promoter','share','group','related_party'],40),['Management / Ownership Field','Value']); _table(doc,[[k,v] for k,v in (add.get('credit_bureau') or {}).items()],['Credit Parameter','Value'])

    # 12
    _page_header(doc,12,PAGE_TITLES[11]); doc.add_paragraph(ai.get('assessment') or 'Justification is based on source-backed business facts, calculated financial metrics and verification exceptions.'); _table(doc,_rows_matching(raw,['associate','subsidiary','group','related'],35),['Group / Associate Evidence','Value'])

    # 13
    _page_header(doc,13,PAGE_TITLES[12]); doc.add_paragraph(ai.get('recommendation') or 'Final recommendation remains subject to policy analysis and authorized credit-officer review.'); _table(doc,[[c.get('rule'),c.get('status'),c.get('reason')] for c in state.get('policy_checks',[])],['Policy Rule','Status','Reason']) if state.get('policy_checks') else doc.add_paragraph('No policy-check result stored for this application.')

    # 14
    _page_header(doc,14,PAGE_TITLES[13]); _table(doc,[['Total Assets',_fmt_money(_field(raw,'total_assets_current'),unit)],['Current Assets',_fmt_money(_field(raw,'total_current_assets_current'),unit)],['Inventory',_fmt_money(_field(raw,'inventories_current'),unit)],['Trade Receivables',_fmt_money(_field(raw,'trade_receivables_current'),unit)],['Cash & Cash Equivalents',_fmt_money(_field(raw,'cash_and_cash_equivalents_current'),unit)],['Equity Share Capital',_fmt_money(_field(raw,'equity_share_capital_current'),unit)],['Other Equity',_fmt_money(_field(raw,'other_equity_current'),unit)],['Non-current Borrowings',_fmt_money(_field(raw,'non_current_borrowings_current'),unit)],['Current Borrowings',_fmt_money(_field(raw,'current_borrowings_current'),unit)],['Current Liabilities',_fmt_money(_field(raw,'total_current_liabilities_current'),unit)]],['Balance Sheet Particular','Latest Raw Value'])

    # 15
    _page_header(doc,15,PAGE_TITLES[14]); _table(doc,[['Revenue from Operations',_fmt_money(_field(raw,'revenue_from_operations_current'),unit)],['Other Income',_fmt_money(_field(raw,'other_income_current'),unit)],['Total Income',_fmt_money(_field(raw,'total_income_current'),unit)],['Total Expenses',_fmt_money(_field(raw,'total_expenses_current'),unit)],['Finance Cost',_fmt_money(_field(raw,'finance_cost_current'),unit)],['Depreciation & Amortisation',_fmt_money(_field(raw,'depreciation_amortisation_current'),unit)],['PBT',_fmt_money(_field(raw,'profit_before_tax_current'),unit)],['PAT',_fmt_money(_field(raw,'profit_after_tax_current'),unit)],['System-calculated EBITDA',_fmt_money(calc.get('ebitda_raw_unit'),unit)]],['Profitability Particular','Value']); _table(doc,[[k,v] for k,v in ratios.items()],['Ratio','Calculated Value'])

    # 16
    _page_header(doc,16,PAGE_TITLES[15]); _table(doc,[['Operating Cash Flow',_fmt_money(_field(raw,'net_cash_from_operating_activities_current'),unit)],['Investing Cash Flow',_fmt_money(_field(raw,'net_cash_from_investing_activities_current'),unit)],['Financing Cash Flow',_fmt_money(_field(raw,'net_cash_from_financing_activities_current'),unit)],['Inventory',_fmt_money(_field(raw,'inventories_current'),unit)],['Trade Receivables',_fmt_money(_field(raw,'trade_receivables_current'),unit)],['Trade Payables',_fmt_money(_field(raw,'trade_payables_current'),unit)]],['Working Capital / Cash Flow','Raw Value']); doc.add_paragraph('Operating-cycle days are calculated only when required raw values and the applicable denominator are available.')

    # 17
    _page_header(doc,17,PAGE_TITLES[16]); doc.add_paragraph(ai.get('cash_flow') or 'Financial commentary is generated from the normalized financial dataset.'); _table(doc,[[k,v] for k,v in ratios.items()],['Key Indicator','Value']); doc.add_paragraph('Any narrative generated by Groq is non-authoritative and cannot overwrite extracted figures, provider verification, policy rules or deterministic calculations.')

    # 18
    _page_header(doc,18,PAGE_TITLES[17]); doc.add_paragraph(ai.get('benchmarking') or 'Industry / peer analysis is displayed only when current evidence or configured external data is available.'); ip=state.get('industry_peer_cache') or {}; _table(doc,[[k,v] for k,v in (ip.get('industry_overview') or {}).items()],['Industry Field','Value'])

    # 19
    _page_header(doc,19,PAGE_TITLES[18]); _table(doc,_rows_matching(raw,['collateral','security','property','valuation','mortgage','insurance','legal','charge'],50),['Collateral / DD Evidence','Value']); _table(doc,[[k,v] for k,v in (add.get('collateral') or {}).items()],['POC Verification Item','Value']); doc.add_paragraph('Synthetic verification is used only where genuine private verification is unavailable in this POC and is explicitly labelled as such.')

    # 20
    _page_header(doc,20,PAGE_TITLES[19]); _h(doc,'Final Credit View',1); doc.add_paragraph(ai.get('recommendation') or 'Final credit decision remains with the authorized human credit officer / sanctioning authority.'); _h(doc,'Source Register',2); _table(doc,[[d.get('original_filename'),d.get('display_name') or d.get('doc_type'),d.get('status')] for d in documents],['Document','Classified Type','Status']); _h(doc,'Verification Audit',2); _table(doc,[['Verification Provider',ver.get('provider')],['Source Kind',ver.get('source_kind')],['Synthetic Fallback Used',ver.get('synthetic')],['Real Provider Error',ver.get('error_from_real_provider')],['Generated CAM ID',state.get('cam_id')]],['Audit Item','Value']); doc.add_paragraph('Officer Review: ____________________   Date: ____________   Sanctioning Authority: ____________________')
    return doc
